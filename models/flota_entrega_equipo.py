import base64
import logging
from urllib.parse import quote as url_quote

from odoo import models, fields, api, _

_logger = logging.getLogger(__name__)


class FlotaEntregaEquipo(models.Model):
    _name = 'flota.entrega.equipo'
    _description = 'Acta de Recepción y Entrega de Equipos'
    _inherit = ['mail.thread', 'mail.activity.mixin']
    _order = 'fecha desc, id desc'

    name = fields.Char(
        string='Nº de Acta',
        required=True,
        copy=False,
        readonly=True,
        default=lambda self: _('Nuevo'),
        index=True,
        tracking=True
    )
    empleado_id = fields.Many2one(
        'flota.empleado',
        string='Empleado',
        required=True,
        ondelete='restrict',
        index=True,
        tracking=True,
        help="Empleado al que se le entrega o recibe el equipo. Al seleccionarlo se autocompletan Ruta, Localidad, Teléfono de flota y Responsable."
    )
    fecha = fields.Date(string='Fecha', default=fields.Date.context_today, required=True, tracking=True)
    estado = fields.Selection([
        ('draft', 'Borrador'),
        ('confirmado', 'Confirmado'),
    ], string='Estado', default='draft', required=True, tracking=True, index=True)

    # --- DATOS ENTREGA (autocompletados desde el Empleado, editables) ---
    tipo_asignacion = fields.Selection([
        ('vendedor', 'Vendedor'),
        ('distribuidor', 'Distribuidor'),
        ('otro', 'Otro'),
    ], string='Distribuidor / Vendedor', tracking=True)
    preparado_por_id = fields.Many2one(
        'res.users', string='Preparado por',
        default=lambda self: self.env.user, tracking=True
    )
    ruta_id = fields.Many2one('flota.ruta', string='Ruta', tracking=True)
    ubicacion_id = fields.Many2one('flota.ubicacion', string='Localidad', tracking=True)
    recibido_por = fields.Char(string='Recibido por (Responsable)', tracking=True)
    telefono_flota = fields.Char(string='Núm. de Teléfono (Flota)', tracking=True)

    # --- DATOS FLOTA RECIBIDA POR TI (equipo antiguo devuelto) ---
    equipo_recibido_modelo = fields.Char(string='Modelo (Flota Recibida)')
    equipo_recibido_serial = fields.Char(string='Serial / IMEI (Flota Recibida)')

    # --- DATOS IMPRESORA ---
    impresora_modelo = fields.Char(string='Modelo (Impresora)')
    impresora_serial = fields.Char(string='Serial (Impresora)')

    # --- DATOS DE EQUIPO NUEVO ENTREGADO ---
    linea_ids = fields.One2many(
        'flota.entrega.equipo.linea', 'entrega_id',
        string='Equipos Entregados'
    )
    cantidad_equipos = fields.Integer(string='Cantidad de Equipos', compute='_compute_cantidad_equipos', store=True)

    # --- FIRMAS (espacio para firma física, sin firma digital) Y POLÍTICA ---
    entregado_por = fields.Char(
        string='Firma Representante IT',
        default=lambda self: self.env.user.name,
        tracking=True,
        help="Nombre del representante de TI que entrega el equipo. Se imprime sobre la línea de firma "
             "'Firma Representante IT' para que la persona firme físicamente el documento."
    )
    recibio_politica = fields.Selection([
        ('si', 'Sí'),
        ('no', 'No'),
    ], string='¿Recibió la Política de Informática?', tracking=True)
    politica_url = fields.Char(
        string='Enlace a la Política de Informática',
        default=lambda self: self.env['ir.config_parameter'].sudo().get_param(
            'gestion_flota_empleados.politica_informatica_url', default=''
        ),
        help="URL de la política de informática. Se usa para generar el código QR del acta si no se sube una "
             "imagen propia. Puede configurarse de forma general en Ajustes Técnicos > Parámetros del Sistema "
             "con la clave 'gestion_flota_empleados.politica_informatica_url'."
    )
    politica_qr_imagen = fields.Image(
        string='Imagen del QR (opcional)',
        max_width=400, max_height=400,
        help="Suba aquí la foto/imagen del código QR de la política de informática tal como debe imprimirse. "
             "Si no sube ninguna imagen, el sistema genera automáticamente un QR a partir del enlace indicado arriba."
    )
    politica_qr_src = fields.Char(string='QR Política (URL interna)', compute='_compute_politica_qr_src')

    # --- CONTROL DOCUMENTAL DEL FORMATO (editable por si cambia la versión oficial) ---
    codigo_formulario = fields.Char(string='Código de Formulario', default='MS-TE-FO-001')
    version_formulario = fields.Char(string='Versión', default='3.00')

    notas = fields.Text(string='Notas')
    company_id = fields.Many2one('res.company', string='Compañía', default=lambda self: self.env.company)

    # --- VISTA PREVIA EN PDF (se actualiza automáticamente al guardar) ---
    vista_previa_pdf = fields.Binary(string='Vista Previa (PDF)', attachment=False, copy=False)
    vista_previa_pdf_filename = fields.Char(string='Nombre de archivo (vista previa)', copy=False)

    @api.depends('linea_ids.cantidad')
    def _compute_cantidad_equipos(self):
        for rec in self:
            rec.cantidad_equipos = sum(rec.linea_ids.mapped('cantidad'))

    @api.depends('politica_url')
    def _compute_politica_qr_src(self):
        for rec in self:
            if rec.politica_url:
                rec.politica_qr_src = '/report/barcode/QR/%s?width=200&height=200' % url_quote(rec.politica_url, safe='')
            else:
                rec.politica_qr_src = False

    @api.onchange('empleado_id')
    def _onchange_empleado_id(self):
        for rec in self:
            if rec.empleado_id:
                rec.ruta_id = rec.empleado_id.ruta_id
                rec.ubicacion_id = rec.empleado_id.ubicacion_id
                rec.telefono_flota = rec.empleado_id.numero_flota
                rec.recibido_por = rec.empleado_id.name
                if rec.empleado_id.ruta_id and rec.empleado_id.ruta_id.tipo in ('vendedor', 'distribuidor'):
                    rec.tipo_asignacion = rec.empleado_id.ruta_id.tipo
                else:
                    rec.tipo_asignacion = 'otro'

    @api.model_create_multi
    def create(self, vals_list):
        for vals in vals_list:
            if vals.get('name', _('Nuevo')) == _('Nuevo'):
                vals['name'] = self.env['ir.sequence'].next_by_code('flota.entrega.equipo') or _('Nuevo')
        records = super().create(vals_list)
        records.with_context(skip_preview_refresh=True)._actualizar_vista_previa_pdf()
        return records

    def write(self, vals):
        res = super().write(vals)
        if not self.env.context.get('skip_preview_refresh'):
            self.with_context(skip_preview_refresh=True)._actualizar_vista_previa_pdf()
        return res

    def _actualizar_vista_previa_pdf(self):
        """Regenera la vista previa en PDF del acta (se ve del lado derecho del formulario).
        Se ejecuta automáticamente al crear/guardar el registro."""
        report = self.env.ref('gestion_flota_empleados.action_report_flota_entrega_equipo', raise_if_not_found=False)
        if not report:
            return
        for rec in self:
            try:
                pdf_content, _report_type = report._render_qweb_pdf(rec.ids)
                rec.vista_previa_pdf = base64.b64encode(pdf_content)
                rec.vista_previa_pdf_filename = '%s.pdf' % (rec.name or 'Acta')
            except Exception:
                _logger.exception("No se pudo generar la vista previa en PDF del acta %s", rec.id)

    def action_actualizar_vista_previa(self):
        self.with_context(skip_preview_refresh=True)._actualizar_vista_previa_pdf()

    def action_confirmar(self):
        self.write({'estado': 'confirmado'})

    def action_restablecer_borrador(self):
        self.write({'estado': 'draft'})

    def action_exportar_excel(self):
        """Exporta el acta de entrega/recepción de equipos a un libro Excel con un diseño limpio y profesional."""
        self.ensure_one()
        import io
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Acta Entrega Equipos"

        NUM_COLS = 7
        thin = Side(style='thin', color='B0B0B0')
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
        header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        title_font = Font(name="Calibri", size=14, bold=True, color="1F2937")
        subtitle_font = Font(name="Calibri", size=9, italic=True, color="6B7280")
        section_fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
        section_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        label_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
        bold_font = Font(name="Calibri", size=10, bold=True)
        normal_font = Font(name="Calibri", size=10)

        def merge_section(row_num, text, fill=section_fill, font=section_font):
            ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=NUM_COLS)
            cell = ws.cell(row=row_num, column=1, value=text)
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(horizontal='center', vertical='center')
            for col in range(1, NUM_COLS + 1):
                ws.cell(row=row_num, column=col).border = border

        def data_row(row_num, label, value):
            ws.merge_cells(start_row=row_num, start_column=2, end_row=row_num, end_column=NUM_COLS)
            lbl_cell = ws.cell(row=row_num, column=1, value=label)
            lbl_cell.font = bold_font
            lbl_cell.fill = label_fill
            val_cell = ws.cell(row=row_num, column=2, value=value)
            val_cell.font = normal_font
            for col in range(1, NUM_COLS + 1):
                ws.cell(row=row_num, column=col).border = border

        ws.merge_cells('A1:G1')
        ws["A1"] = f"ACTA DE RECEPCIÓN Y ENTREGA DE EQUIPOS — {self.name}"
        ws["A1"].font = title_font
        ws["A1"].alignment = Alignment(horizontal='center')
        ws.merge_cells('A2:G2')
        ws["A2"] = f"Código: {self.codigo_formulario or ''}   |   Versión: {self.version_formulario or ''}   |   Fecha: {self.fecha}   |   Empleado: {self.empleado_id.name or ''}"
        ws["A2"].font = subtitle_font
        ws["A2"].alignment = Alignment(horizontal='center')

        row = 4
        merge_section(row, "DATOS ENTREGA")
        row += 1
        datos_entrega = [
            ("Distribuidor/Vendedor", dict(self._fields['tipo_asignacion'].selection).get(self.tipo_asignacion, '') if self.tipo_asignacion else ''),
            ("Preparado por", self.preparado_por_id.name or ''),
            ("Ruta", self.ruta_id.name or ''),
            ("Localidad", self.ubicacion_id.name or ''),
            ("Recibido por (responsable)", self.recibido_por or ''),
            ("Núm. de teléfono (flota)", self.telefono_flota or ''),
        ]
        for label, value in datos_entrega:
            data_row(row, label, value)
            row += 1

        row += 1
        merge_section(row, "DATOS FLOTA RECIBIDA POR TI")
        row += 1
        for label, value in [("Modelo", self.equipo_recibido_modelo or 'N/A'), ("Serial/IMEI", self.equipo_recibido_serial or 'N/A')]:
            data_row(row, label, value)
            row += 1

        row += 1
        merge_section(row, "DATOS IMPRESORA")
        row += 1
        for label, value in [("Modelo", self.impresora_modelo or 'N/A'), ("Serial", self.impresora_serial or 'N/A')]:
            data_row(row, label, value)
            row += 1

        row += 1
        merge_section(row, "DATOS DE EQUIPO NUEVO ENTREGADO")
        row += 1
        headers_eq = ["TIPO DE EQUIPO", "MARCA", "MODELO", "CANT.", "IMEI / SERIAL", "ESTADO", "OBSERVACIONES"]
        for col_num, h in enumerate(headers_eq, 1):
            cell = ws.cell(row=row, column=col_num, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = border
        row += 1
        for line in self.linea_ids:
            values = [
                dict(line._fields['tipo_equipo'].selection).get(line.tipo_equipo, ''),
                line.marca or '',
                line.modelo or '',
                line.cantidad,
                line.imei_serial or '',
                dict(line._fields['estado_equipo'].selection).get(line.estado_equipo, ''),
                line.observaciones or '',
            ]
            for col_num, value in enumerate(values, 1):
                cell = ws.cell(row=row, column=col_num, value=value)
                cell.font = normal_font
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                cell.border = border
            row += 1

        row += 2
        merge_section(row, "FIRMAS Y POLÍTICA DE INFORMÁTICA")
        row += 1
        data_row(row, "Firma Representante IT", self.entregado_por or '')
        row += 1
        data_row(row, "Recibido Por", self.recibido_por or '')
        row += 1
        data_row(row, "¿Recibió la Política de Informática?", dict(self._fields['recibio_politica'].selection).get(self.recibio_politica, ''))
        row += 1
        if self.politica_url:
            data_row(row, "Enlace Política de Informática", self.politica_url)
            row += 1

        row += 1
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=NUM_COLS)
        legal_cell = ws.cell(row=row, column=1, value=(
            "Mediante la firma de este documento, comprendo y asumo la responsabilidad que me confiere la asignación "
            "de los equipos aquí detallados y entiendo que la violación a cualquiera de las directivas establecidas "
            "en la Política de Informática, la cual he recibido, leído y entendido, puede conllevar a que la empresa "
            "revoque mis privilegios y tome acciones disciplinarias y/o legales de acuerdo con lo establecido en "
            "dicha política."
        ))
        legal_cell.font = Font(name="Calibri", size=9, italic=True, color="4B5563")
        legal_cell.alignment = Alignment(horizontal='justify', vertical='top', wrap_text=True)
        ws.row_dimensions[row].height = 60

        column_widths = [22, 20, 20, 8, 20, 14, 30]
        for idx, width in enumerate(column_widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = width

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        file_data = base64.b64encode(output.read())
        output.close()

        attachment = self.env['ir.attachment'].create({
            'name': f'{self.name} - {self.empleado_id.name}.xlsx',
            'datas': file_data,
            'res_model': self._name,
            'res_id': self.id,
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        })

        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/{attachment.id}?download=true',
            'target': 'self',
        }


class FlotaEntregaEquipoLinea(models.Model):
    _name = 'flota.entrega.equipo.linea'
    _description = 'Línea de Equipo Entregado'
    _order = 'id asc'

    entrega_id = fields.Many2one('flota.entrega.equipo', string='Acta de Entrega', required=True, ondelete='cascade', index=True)
    tipo_equipo = fields.Selection([
        ('celular', 'Celular'),
        ('tablet', 'Tablet'),
        ('laptop', 'Laptop'),
        ('impresora', 'Impresora'),
        ('accesorio', 'Accesorio'),
        ('otro', 'Otro'),
    ], string='Tipo de Equipo', required=True, default='celular')
    marca = fields.Char(string='Marca')
    modelo = fields.Char(string='Modelo')
    cantidad = fields.Integer(string='Cant.', default=1, required=True)
    imei_serial = fields.Char(string='IMEI / Serial')
    estado_equipo = fields.Selection([
        ('nueva', 'Nueva'),
        ('usada', 'Usada'),
    ], string='Estado', default='nueva')
    observaciones = fields.Char(
        string='Estado / Observaciones',
        help="Ej. Protector de pantalla roto, cámara cristal roto, se friza a veces."
    )
