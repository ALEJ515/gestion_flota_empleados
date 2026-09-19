import json
import logging
import base64
import io
import re
from odoo import models, fields, api, _
from odoo.exceptions import UserError, ValidationError

from .flota_factura_claro_parser import (
    CLARO_COLUMN_NAMES,
    build_empleado_phone_map,
    extract_claro_phone_row,
    fix_claro_pdf_line,
    map_claro_columns,
    normalize_phone,
    parse_money_token,
)

_logger = logging.getLogger(__name__)


def _normalize_phone(phone_str):
    return normalize_phone(phone_str)


def _build_empleado_phone_map(env):
    return build_empleado_phone_map(env)


def _fix_claro_pdf_line(line_str):
    return fix_claro_pdf_line(line_str)


class FlotaFacturaConciliacion(models.Model):
    _name = 'flota.factura.conciliacion'
    _description = 'Conciliación Financiera de Factura de Flota'
    _inherit = ['mail.thread', 'mail.activity.mixin']
    _order = 'periodo desc, id desc'

    name = fields.Char(string='Referencia / Folio', required=True, copy=False, default=lambda self: _('Nuevo'), index=True, tracking=True, help="Número de referencia o folio único asignado a la factura de flota.")
    proveedor = fields.Char(string='Proveedor', default='Claro Dominicana', required=True, tracking=True, help="Empresa prestadora del servicio telefónico corporativo (Claro Dominicana).")
    periodo = fields.Char(string='Periodo / Mes (AAAA-MM)', required=True, default=lambda self: fields.Date.today().strftime('%Y-%m'), index=True, tracking=True, help="Año y mes fiscal correspondiente a esta conciliación (Ej. 2026-08).")
    fecha_factura = fields.Date(string='Fecha de Factura', default=fields.Date.context_today, required=True, tracking=True, help="Fecha exacta de emisión o facturación impresa en el documento de Claro.")
    fecha_subida = fields.Datetime(string='Fecha de Subida', default=fields.Datetime.now, readonly=True, help="Fecha y hora de registro y subida del PDF a Odoo.")
    
    archivo_pdf = fields.Binary(string='Adjuntar PDF Factura Claro', attachment=True, help="Documento digital PDF de la factura corporativa Claro para extracción nativa.")
    pdf_filename = fields.Char(string='Nombre del Archivo PDF')

    fecha_factura_str = fields.Char(string='Fecha / Mes Formateado', compute='_compute_fecha_factura_str', store=True, help="Fecha y mes en formato legible en español (ej. 15 de Agosto de 2026).")

    renta_mensual = fields.Monetary(string='Renta Mensual', currency_field='currency_id', default=0.0, tracking=True, help="Monto total acumulado por renta fija de planes corporativos de voz y data.")
    renta_otros_servicios = fields.Monetary(string='Renta Otros Servicios', currency_field='currency_id', default=0.0, tracking=True, help="Servicios adicionales y paquetes complementarios contratados en la cuenta.")
    uso_data_movil = fields.Monetary(string='Uso Data Móvil', currency_field='currency_id', default=0.0, tracking=True, help="Consumo excedente de datos móviles no incluidos en los planes fijos.")
    llamadas_roaming = fields.Monetary(string='Llamadas Roaming', currency_field='currency_id', default=0.0, tracking=True, help="Consumos por llamadas de Larga Distancia Internacional (LDI) o Roaming.")
    otros_cargos_creditos = fields.Monetary(string='Otros cargos, créditos o descuentos', currency_field='currency_id', default=0.0, tracking=True, help="Cargos extraordinarios, notas de crédito o descuentos corporativos globales a nivel de cuenta.")

    subtotal = fields.Monetary(string='Subtotal Factura', compute='_compute_totales_factura', store=True, currency_field='currency_id', tracking=True, help="Base imponible de la factura antes de aplicar los impuestos de ley (RD$).")
    subtotal_factura = fields.Monetary(string='Subtotal Factura', compute='_compute_totales_factura', store=True, currency_field='currency_id', tracking=True, help="Base imponible de la factura antes de aplicar los impuestos de ley (RD$).")
    base_gravable_itbis = fields.Monetary(string='Base Gravable ITBIS', compute='_compute_totales_factura', store=True, currency_field='currency_id', help="Monto total gravado con el 18% de ITBIS.")
    base_gravable_cdt = fields.Monetary(string='Base Gravable CDT', compute='_compute_totales_factura', store=True, currency_field='currency_id', help="Monto total gravado con el 2% de CDT (Subtotal menos Conceptos Excluidos).")
    base_gravable_isc = fields.Monetary(string='Base Gravable ISC', compute='_compute_totales_factura', store=True, currency_field='currency_id', help="Monto total gravado con el 10% de ISC.")
    ajustes_excluidos_cdt = fields.Monetary(string='Conceptos Excluidos CDT', compute='_compute_totales_factura', store=True, readonly=False, currency_field='currency_id', help="Conceptos/Ajustes que NO gravan CDT (ej. Cargo por Pago Atrasado / Mora).")

    itbis_monto = fields.Monetary(string='ITBIS (18%)', compute='_compute_totales_factura', store=True, readonly=False, currency_field='currency_id', tracking=True, help="Impuesto a la Transferencia de Bienes Industrializados y Servicios (18%). Se usa el monto impreso literalmente en el PDF; si no se pudo extraer, se calcula como 18% del Subtotal.")
    itbis_pdf_extraido = fields.Monetary(string='ITBIS Extraído del PDF', currency_field='currency_id', default=0.0, help="Monto de ITBIS impreso literalmente en la carátula del PDF de Claro (fuente oficial, evita diferencias por redondeo interno de Claro al calcular por línea).")
    cdt_monto = fields.Monetary(string='CDT informado (2%)', currency_field='currency_id', default=0.0, tracking=True, help="Monto CDT usado para obtener la base gravable mediante CDT / 2%.")
    isc_monto = fields.Monetary(string='ISC (10%)', compute='_compute_totales_factura', store=True, readonly=False, currency_field='currency_id', tracking=True, help="Impuesto Selectivo al Consumo de Telecomunicaciones (10%). Se usa el monto impreso literalmente en el PDF; si no se pudo extraer, se calcula como 10% del Subtotal.")
    isc_pdf_extraido = fields.Monetary(string='ISC Extraído del PDF', currency_field='currency_id', default=0.0, help="Monto de ISC impreso literalmente en la carátula del PDF de Claro (fuente oficial, evita diferencias por redondeo interno de Claro al calcular por línea).")
    total_mes = fields.Monetary(string='Total del Mes', compute='_compute_totales_factura', store=True, currency_field='currency_id', tracking=True, help="Gran Total final de la factura Claro en RD$ a pagar (Subtotal + Impuestos).")
    total_factura_pdf = fields.Monetary(string='Total Factura Claro (PDF)', currency_field='currency_id', default=0.0, tracking=True, help="Monto 'Total del Mes' impreso literalmente en la carátula del PDF de Claro. Se usa como referencia oficial para validar la conciliación.")
    formula_calculo_html = fields.Html(string='Cómo se calculó el Total del Mes', compute='_compute_totales_factura', sanitize=False, help="Detalle paso a paso, con los montos actuales de esta factura, de cómo el sistema obtiene el Total del Mes.")

    currency_id = fields.Many2one('res.currency', string='Moneda', default=lambda self: self.env.company.currency_id)
    estado = fields.Selection([
        ('draft', 'Borrador'),
        ('procesando', 'Procesando'),
        ('conciliado', 'Conciliado'),
        ('error', 'Con Errores / Desviaciones')
    ], string='Estado', default='draft', required=True, index=True, tracking=True)

    notes = fields.Text(string='Observaciones y Notas')

    linea_ids = fields.One2many('flota.factura.linea', 'conciliacion_id', string='Desglose por Empleado / Número')
    resumen_depto_ids = fields.One2many('flota.factura.departamento.resumen', 'conciliacion_id', string='Resumen por Departamento')
    concepto_ids = fields.One2many('flota.factura.concepto', 'conciliacion_id', string='Conceptos y Ajustes de Factura')

    count_lineas = fields.Integer(string='Total Líneas PDF', compute='_compute_kpis', store=True, help='Cantidad de líneas telefónicas extraídas del PDF de la factura.')
    count_lineas_registradas = fields.Integer(string='Total Líneas Registradas', compute='_compute_count_lineas_registradas', store=False, help='Cantidad total de Empleados / Números de Flota registrados actualmente en el módulo (Perfil de Empleados y Flotas), incluyendo activos e inactivos. Se calcula en tiempo real (no almacenado) para reflejar altas y bajas de empleados de forma inmediata.')
    count_excesos = fields.Integer(string='Líneas con Exceso', compute='_compute_kpis', store=True)
    monto_excesos = fields.Monetary(string='Monto Total Excesos', compute='_compute_kpis', store=True, currency_field='currency_id')

    total_lineas_sum = fields.Monetary(
        string='Suma Consumo Líneas (RD$)',
        compute='_compute_kpis',
        store=True,
        currency_field='currency_id',
        tracking=True
    )
    estado_cuadre = fields.Selection([
        ('cuadrado', 'Cuadrado Exacto'),
        ('desviacion', 'Con Diferencia')
    ], string='Estado de Conciliación', compute='_compute_kpis', store=True, tracking=True)

    banner_conciliacion_html = fields.Html(
        string='Observación de la Conciliación',
        compute='_compute_kpis',
        store=True
    )

    @api.model_create_multi
    def create(self, vals_list):
        for vals in vals_list:
            if vals.get('name', _('Nuevo')) == _('Nuevo'):
                vals['name'] = self.env['ir.sequence'].next_by_code('flota.factura.conciliacion') or _('FAC-CLARO-%s') % fields.Date.today()
        return super(FlotaFacturaConciliacion, self).create(vals_list)

    def write(self, vals):
        # Se detecta el cambio a 'conciliado' sin importar si proviene de un botón o de un
        # clic directo en el statusbar, para no depender de action_marcar_conciliado como único
        # punto de entrada y así poder simplificar los botones del header sin perder este efecto.
        pasa_a_conciliado = vals.get('estado') == 'conciliado' and any(
            rec.estado != 'conciliado' for rec in self
        )
        res = super(FlotaFacturaConciliacion, self).write(vals)
        if pasa_a_conciliado:
            self.action_generar_resumen_departamentos()
            self.message_post(body=_("Factura de Flota marcada como Conciliada correctamente."))
        elif vals.get('estado') == 'draft':
            self.message_post(body=_("El estado de la conciliación fue restablecido a <b>Borrador</b>."))
        elif vals.get('estado') == 'procesando':
            self.message_post(body=_("El estado de la conciliación fue cambiado a <b>Procesando</b>."))
        return res

    @api.depends('fecha_factura')
    def _compute_fecha_factura_str(self):
        meses = {
            1: 'Enero', 2: 'Febrero', 3: 'Marzo', 4: 'Abril',
            5: 'Mayo', 6: 'Junio', 7: 'Julio', 8: 'Agosto',
            9: 'Septiembre', 10: 'Octubre', 11: 'Noviembre', 12: 'Diciembre'
        }
        for rec in self:
            if rec.fecha_factura:
                m_name = meses.get(rec.fecha_factura.month, '')
                rec.fecha_factura_str = f"{rec.fecha_factura.day} de {m_name} de {rec.fecha_factura.year}"
            else:
                rec.fecha_factura_str = ''

    def action_set_draft(self):
        self.ensure_one()
        self.write({'estado': 'draft'})

    def action_set_procesando(self):
        self.ensure_one()
        self.write({'estado': 'procesando'})

    @api.depends(
        'renta_mensual', 'renta_otros_servicios', 'uso_data_movil', 'llamadas_roaming',
        'otros_cargos_creditos', 'cdt_monto', 'itbis_pdf_extraido', 'isc_pdf_extraido'
    )
    def _compute_totales_factura(self):
        for rec in self:
            # Los créditos y descuentos se almacenan negativos y se suman al subtotal.
            sub = (
                rec.renta_mensual
                + rec.renta_otros_servicios
                + rec.uso_data_movil
                + rec.llamadas_roaming
                + rec.otros_cargos_creditos
            )
            rec.subtotal = sub
            rec.subtotal_factura = sub
            rec.base_gravable_itbis = sub
            rec.base_gravable_isc = sub
            rec.base_gravable_cdt = rec.cdt_monto / 0.02 if rec.cdt_monto else 0.0
            rec.ajustes_excluidos_cdt = sub - rec.base_gravable_cdt
            # Claro calcula ITBIS/ISC línea por línea (con redondeo por empleado) y luego suma,
            # por lo que el monto impreso en el PDF puede diferir en centavos del 18%/10% del Subtotal.
            # Se usa el monto extraído literal del PDF como fuente oficial; si no se pudo extraer,
            # se recurre a la fórmula (18%/10% del Subtotal) como respaldo.
            rec.itbis_monto = rec.itbis_pdf_extraido if rec.itbis_pdf_extraido else sub * 0.18
            rec.isc_monto = rec.isc_pdf_extraido if rec.isc_pdf_extraido else sub * 0.10
            rec.total_mes = sub + rec.itbis_monto + rec.isc_monto + rec.cdt_monto

            itbis_origen = _('extraído del PDF') if rec.itbis_pdf_extraido else _('calculado como 18%s del Subtotal (no se encontró en el PDF)') % '%'
            isc_origen = _('extraído del PDF') if rec.isc_pdf_extraido else _('calculado como 10%s del Subtotal (no se encontró en el PDF)') % '%'
            rec.formula_calculo_html = (
                '<details class="text-muted small">'
                '<summary style="cursor:pointer;">%s</summary>'
                '<table class="table table-sm mb-0 mt-2" style="max-width:520px;">'
                '<tbody>'
                '<tr><td>Renta Mensual Planes</td><td class="text-end">RD$%s</td></tr>'
                '<tr><td>Renta Otros Servicios</td><td class="text-end">RD$%s</td></tr>'
                '<tr><td>Uso Data Móvil</td><td class="text-end">RD$%s</td></tr>'
                '<tr><td>Llamadas Roaming / LD</td><td class="text-end">RD$%s</td></tr>'
                '<tr><td>Otros cargos, créditos o descuentos</td><td class="text-end">RD$%s</td></tr>'
                '<tr class="fw-bold border-top"><td>(=) Subtotal Factura</td><td class="text-end">RD$%s</td></tr>'
                '<tr><td colspan="2" class="pt-3"><em>Base Gravable CDT</em> = CDT informado (2%%) ÷ 0.02</td></tr>'
                '<tr><td>CDT informado (2%%)</td><td class="text-end">RD$%s</td></tr>'
                '<tr><td>(=) Base Gravable CDT</td><td class="text-end">RD$%s</td></tr>'
                '<tr><td colspan="2" class="pt-3"><em>Impuestos</em> (usan el monto impreso en el PDF; si falta, se calculan sobre el Subtotal)</td></tr>'
                '<tr><td>ITBIS (18%%) &mdash; %s</td><td class="text-end">RD$%s</td></tr>'
                '<tr><td>ISC (10%%) &mdash; %s</td><td class="text-end">RD$%s</td></tr>'
                '<tr><td>CDT (2%%)</td><td class="text-end">RD$%s</td></tr>'
                '<tr class="fw-bold border-top"><td>(=) TOTAL DEL MES</td><td class="text-end">RD$%s</td></tr>'
                '</tbody>'
                '</table>'
                '<div class="mt-2">Fórmula: Total del Mes = Subtotal + ITBIS + ISC + CDT informado.</div>'
                '</details>'
            ) % (
                _('Ver cómo se calculó el Total del Mes ▾'),
                f"{rec.renta_mensual:,.2f}",
                f"{rec.renta_otros_servicios:,.2f}",
                f"{rec.uso_data_movil:,.2f}",
                f"{rec.llamadas_roaming:,.2f}",
                f"{rec.otros_cargos_creditos:,.2f}",
                f"{sub:,.2f}",
                f"{rec.cdt_monto:,.2f}",
                f"{rec.base_gravable_cdt:,.2f}",
                itbis_origen,
                f"{rec.itbis_monto:,.2f}",
                isc_origen,
                f"{rec.isc_monto:,.2f}",
                f"{rec.cdt_monto:,.2f}",
                f"{rec.total_mes:,.2f}",
            )

    @api.depends('linea_ids', 'linea_ids.total', 'linea_ids.total_linea', 'linea_ids.estado_linea', 'linea_ids.empleado_id', 'linea_ids.empleado_id.es_nuevo_auto', 'linea_ids.uso_local_data_movil', 'linea_ids.monto_roaming', 'total_mes', 'total_factura_pdf')
    def _compute_kpis(self):
        for rec in self:
            rec.count_lineas = len(rec.linea_ids)
            excesos = rec.linea_ids.filtered(lambda l: l.estado_linea in ['exceso_data', 'exceso_roaming'])
            rec.count_excesos = len(excesos)
            rec.monto_excesos = sum(excesos.mapped(lambda l: l.uso_local_data_movil + l.monto_roaming))

            tot_lineas = sum(rec.linea_ids.mapped('total'))
            rec.total_lineas_sum = tot_lineas
            # La conciliación se valida contra el Total de la Factura Claro impreso en el PDF, no contra la suma de líneas de empleados.
            diff = rec.total_factura_pdf - rec.total_mes

            if abs(diff) < 0.01:
                rec.estado_cuadre = 'cuadrado'
                rec.banner_conciliacion_html = False
            else:
                rec.estado_cuadre = 'desviacion'
                rec.banner_conciliacion_html = (
                    '<div class="alert alert-warning d-flex align-items-center mb-3 shadow-sm" role="alert">'
                    '<i class="fa fa-exclamation-triangle fs-4 me-2"></i>'
                    '<div>'
                    '<strong>Observación de la Conciliación:</strong><br/>'
                    'El Total de la Factura Claro impreso en el PDF es de <strong>RD$%s</strong>.<br/>'
                    'El Total del Mes calculado por el sistema es de <strong>RD$%s</strong>, sumando el Subtotal (<strong>RD$%s</strong>) más los impuestos '
                    '(ITBIS 18%%: <strong>RD$%s</strong>, ISC 10%%: <strong>RD$%s</strong> y CDT 2%%: <strong>RD$%s</strong>).<br/>'
                    'Existe una diferencia de <strong>RD$%s</strong> entre el monto de la factura y el monto calculado por el sistema. '
                    'Revise los valores de Renta, Otros cargos/créditos y CDT capturados de la factura.'
                    '</div>'
                    '</div>'
                ) % (
                    f"{rec.total_factura_pdf:,.2f}",
                    f"{rec.total_mes:,.2f}",
                    f"{rec.subtotal:,.2f}",
                    f"{rec.itbis_monto:,.2f}",
                    f"{rec.isc_monto:,.2f}",
                    f"{rec.cdt_monto:,.2f}",
                    f"{abs(diff):,.2f}"
                )

    def _compute_count_lineas_registradas(self):
        """Campo no almacenado: refleja en tiempo real el total de Empleados / Números
        de Flota registrados en el módulo (Perfil de Empleados y Flotas), incluyendo
        activos e inactivos. Al no tener store=True, Odoo lo recalcula cada vez que se
        lee/abre el registro, por lo que altas, bajas o eliminaciones de empleados se
        reflejan de inmediato sin necesidad de re-procesar cada factura.
        """
        total_registrados = self.env['flota.empleado'].with_context(active_test=False).search_count([])
        for rec in self:
            rec.count_lineas_registradas = total_registrados

    def action_generar_resumen_departamentos(self):
        """ Agrupa y consolida el gasto por Departamento, impacta historial y sincroniza estado (Activo/Inactivo) de Empleados """
        for rec in self:
            rec.resumen_depto_ids.unlink()
            dept_totals = {}
            seen_emp_ids = set()
            seen_phones = set()

            for linea in rec.linea_ids:
                if linea.numero_flota:
                    norm = _normalize_phone(linea.numero_flota)
                    if norm:
                        seen_phones.add(norm)
                        if len(norm) >= 10:
                            seen_phones.add(norm[-10:])
                        if len(norm) >= 7:
                            seen_phones.add(norm[-7:])

                if linea.empleado_id:
                    seen_emp_ids.add(linea.empleado_id.id)

                dept_id = linea.departamento_id.id if linea.departamento_id else 0

                if dept_id not in dept_totals:
                    # La ubicación del consolidado prioriza la configurada en el Departamento
                    # (flota.departamento.ubicacion_id); si el departamento no tiene ubicación
                    # asignada, se usa como respaldo la ubicación individual del empleado.
                    ubicacion_dept = linea.departamento_id.ubicacion_id if linea.departamento_id else False
                    ubicacion_final = ubicacion_dept or linea.ubicacion_id
                    dept_totals[dept_id] = {
                        'departamento_id': linea.departamento_id.id if linea.departamento_id else False,
                        'ubicacion_id': ubicacion_final.id if ubicacion_final else False,
                        'cantidad_empleados': 0,
                        'monto_subtotal': 0.0,
                        'monto_total': 0.0
                    }
                dept_totals[dept_id]['cantidad_empleados'] += 1
                dept_totals[dept_id]['monto_subtotal'] += linea.subtotal_linea
                dept_totals[dept_id]['monto_total'] += linea.total_linea

            if seen_emp_ids:
                self.env['flota.empleado'].browse(list(seen_emp_ids))._update_facturacion_stats()

            # Sincronización automática de estado de Empleados (Activo vs Inactivo).
            # Se procesa en lote (bulk write + log batch) en vez de un write()/message_post()
            # por empleado: con cientos de empleados, hacerlo uno por uno es la causa principal
            # de que la extracción del PDF se quede "cargando" (cada message_post individual es
            # una operación pesada de mail.thread que puede sumar minutos en facturas grandes).
            all_emps = self.env['flota.empleado'].with_context(active_test=False).search([])
            to_activate_ids = []
            to_deactivate_ids = []
            log_bodies = {}
            for emp in all_emps:
                if not emp.numero_flota:
                    continue
                emp_norm = _normalize_phone(emp.numero_flota)
                if not emp_norm:
                    continue

                is_present = (
                    emp.id in seen_emp_ids or
                    emp_norm in seen_phones or
                    (len(emp_norm) >= 10 and emp_norm[-10:] in seen_phones) or
                    (len(emp_norm) >= 7 and emp_norm[-7:] in seen_phones)
                )

                if is_present:
                    if emp.estado != 'active' or not emp.en_ultima_factura:
                        to_activate_ids.append(emp.id)
                        log_bodies[emp.id] = _("Estado del empleado actualizado a <b>Activo</b> al ser detectado en la conciliación de Claro (%s).") % rec.periodo
                else:
                    if emp.estado != 'inactive' or emp.en_ultima_factura:
                        to_deactivate_ids.append(emp.id)
                        log_bodies[emp.id] = _("<b>Revisión de Flota:</b> Empleado NO detectado en la factura Claro del periodo (%s). Marcado como Faltante / Inactivo.") % rec.periodo

            Empleado = self.env['flota.empleado']
            if to_activate_ids:
                Empleado.browse(to_activate_ids).with_context(tracking_disable=True).write({'estado': 'active', 'en_ultima_factura': True})
            if to_deactivate_ids:
                Empleado.browse(to_deactivate_ids).with_context(tracking_disable=True).write({'estado': 'inactive', 'en_ultima_factura': False})
            if log_bodies:
                Empleado.browse(list(log_bodies.keys()))._message_log_batch(bodies=log_bodies)

            resumen_vals = []
            tot_gral = rec.total_mes
            for d_id, data in dept_totals.items():
                pct = (data['monto_total'] / tot_gral) if tot_gral else 0.0
                resumen_vals.append((0, 0, {
                    'conciliacion_id': rec.id,
                    'departamento_id': data['departamento_id'],
                    'ubicacion_id': data['ubicacion_id'],
                    'cantidad_empleados': data['cantidad_empleados'],
                    'monto_subtotal': data['monto_subtotal'],
                    'monto_total': data['monto_total'],
                    'porcentaje_gasto': pct,
                    'currency_id': rec.currency_id.id
                }))
                
                if data['departamento_id']:
                    dept_rec = self.env['flota.departamento'].browse(data['departamento_id'])
                    if dept_rec.exists():
                        new_vals = {
                            'ultima_facturacion_monto': data['monto_total'],
                            'ultima_facturacion_periodo': str(rec.periodo or '')
                        }
                        if dept_rec.ultima_facturacion_periodo and dept_rec.ultima_facturacion_periodo != str(rec.periodo or ''):
                            new_vals['penultima_facturacion_monto'] = dept_rec.ultima_facturacion_monto
                            new_vals['penultima_facturacion_periodo'] = dept_rec.ultima_facturacion_periodo
                        dept_rec.write(new_vals)

            rec.write({'resumen_depto_ids': resumen_vals})

    def action_marcar_conciliado(self):
        self.ensure_one()
        self.write({'estado': 'conciliado'})

    def _extract_pdf_text_native(self, pdf_bytes):
        """Devuelve (texto, extractor) donde extractor es 'pypdf' o 'pdfplumber'.

        IMPORTANTE: pypdf/PyPDF2 extrae correctamente el texto de la tabla
        "Resumen Factura del Mes por Número" de Claro, pero reordena los
        valores numéricos de cada línea de forma distinta al orden visual
        real de las columnas (confirmado estadísticamente sobre cientos de
        líneas de varias facturas reales). pdfplumber, en cambio, sí preserva
        el orden visual correcto. Por eso se debe saber con qué extractor se
        obtuvo el texto, para poder corregir el orden de columnas en
        map_claro_columns cuando el texto viene de pypdf.
        """
        text = ""
        extractor = 'pypdf'
        # 1. Intentar primero con pypdf/PyPDF2 (Rápido, ultra liviano en RAM)
        try:
            import io
            try:
                from pypdf import PdfReader
                reader = PdfReader(io.BytesIO(pdf_bytes))
                for page in reader.pages:
                    text += (page.extract_text() or "") + "\n"
            except Exception:
                from PyPDF2 import PdfFileReader
                reader = PdfFileReader(io.BytesIO(pdf_bytes))
                for i in range(reader.getNumPages()):
                    text += (reader.getPage(i).extractText() or "") + "\n"
        except Exception as e:
            _logger.warning("Error extrayendo con PyPDF/PyPDF2: %s", str(e))

        if text and len(text.strip()) >= 100:
            return text, extractor

        # 2. Fallback a pdfplumber si pypdf no extrajo texto completo
        try:
            import io
            import pdfplumber
            with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
                for page in pdf.pages:
                    text += (page.extract_text() or "") + "\n"
            extractor = 'pdfplumber'
        except Exception as e2:
            _logger.warning("Error extrayendo con pdfplumber: %s", str(e2))

        return text, extractor

    # Límite de tamaño de archivo PDF para proteger la memoria del servidor: un PDF
    # inusualmente pesado (ej. escaneado como imagen) puede forzar el uso del extractor
    # pdfplumber (más costoso en RAM) sobre un archivo grande. Las facturas reales de
    # Claro pesan típicamente unos cientos de KB, muy por debajo de este límite.
    MAX_PDF_SIZE_BYTES = 2 * 1024 * 1024  # 2 MB

    def action_procesar_pdf_nativo(self):
        """ Extrae y concilia la factura PDF directamente en Odoo """
        self.ensure_one()
        if not self.archivo_pdf:
            raise UserError(_('Por favor adjunte el archivo PDF de la Factura de Claro antes de procesar.'))

        pdf_bytes = base64.b64decode(self.archivo_pdf)

        if len(pdf_bytes) > self.MAX_PDF_SIZE_BYTES:
            raise UserError(_(
                'El archivo PDF adjunto pesa %s MB, superando el límite permitido de %s MB. '
                'Adjunte una versión más liviana del PDF (evite escaneos/imágenes; use el PDF digital original de Claro).'
            ) % (f"{len(pdf_bytes) / (1024 * 1024):.2f}", self.MAX_PDF_SIZE_BYTES // (1024 * 1024)))

        pdf_text, pdf_extractor = self._extract_pdf_text_native(pdf_bytes)

        if not pdf_text or len(pdf_text.strip()) < 20:
            raise UserError(_('No se pudo extraer texto del archivo PDF adjunto.'))

        # Extracción de Fecha de Factura desde el PDF si está presente.
        # Claro imprime la fecha como "Fecha de Factura: Agosto 13,2026" (mes en texto en español),
        # aunque también se soporta el formato numérico DD/MM/AAAA por compatibilidad.
        import datetime
        MESES_ES = {
            'enero': 1, 'febrero': 2, 'marzo': 3, 'abril': 4, 'mayo': 5, 'junio': 6,
            'julio': 7, 'agosto': 8, 'septiembre': 9, 'setiembre': 9, 'octubre': 10,
            'noviembre': 11, 'diciembre': 12,
        }
        fecha_extraida = None
        m_fecha_txt = re.search(
            r'Fecha\s+de\s+Factura\s*:?\s*([A-Za-zÁÉÍÓÚáéíóú]+)\s+([0-9]{1,2})\s*,\s*([0-9]{4})',
            pdf_text, re.IGNORECASE
        )
        if m_fecha_txt:
            mes_nombre = m_fecha_txt.group(1).strip().lower()
            mes_num = MESES_ES.get(mes_nombre)
            if mes_num:
                try:
                    fecha_extraida = datetime.date(int(m_fecha_txt.group(3)), mes_num, int(m_fecha_txt.group(2)))
                except Exception:
                    fecha_extraida = None
        if not fecha_extraida:
            m_fecha = re.search(r'Fecha\s*(?:de\s*factura|facturaci[oó]n|emisi[oó]n)?:?\s*([0-9]{1,2})[/-]([0-9]{1,2})[/-]([0-9]{2,4})', pdf_text, re.IGNORECASE)
            if m_fecha:
                d, m, y = int(m_fecha.group(1)), int(m_fecha.group(2)), int(m_fecha.group(3))
                if y < 100:
                    y += 2000
                try:
                    fecha_extraida = datetime.date(y, m, d)
                except Exception:
                    fecha_extraida = None
        if fecha_extraida:
            self.fecha_factura = fecha_extraida
            # El Periodo (AAAA-MM) debe reflejar el mes fiscal real de la factura, no la fecha
            # en que se subió el PDF a Odoo, para no romper el historial de "última/penúltima factura".
            self.periodo = fecha_extraida.strftime('%Y-%m')

        renta_m = 0.0
        renta_o = 0.0
        data_m = 0.0
        roam_m = 0.0
        cred_m = 0.0
        cdt_m = 0.0
        itbis_m = 0.0
        isc_m = 0.0

        m_renta = re.search(r'Renta\s+mensual\s+([0-9,]+\.[0-9]{2})', pdf_text, re.IGNORECASE)
        if m_renta:
            renta_m = float(m_renta.group(1).replace(',', ''))

        m_renta_o = re.search(r'Renta\s+otros\s+servicios\s+([0-9,]+\.[0-9]{2})', pdf_text, re.IGNORECASE)
        if m_renta_o:
            renta_o = float(m_renta_o.group(1).replace(',', ''))

        m_data = re.search(r'Uso\s+servicios\s+Data\s+M[oó]vil\s+([0-9,]+\.[0-9]{2})', pdf_text, re.IGNORECASE)
        if m_data:
            data_m = float(m_data.group(1).replace(',', ''))

        m_roam = re.search(r'Llamadas\s+(?:Roaming|larga\s+distancia)\s+([0-9,]+\.[0-9]{2})', pdf_text, re.IGNORECASE)
        if m_roam:
            roam_m = float(m_roam.group(1).replace(',', ''))

        m_cred = re.search(r'Otros\s+cargos,?\s+cr[eé]ditos[^\n]*?([0-9,]+\.[0-9]{2}\s*(?:CR)?)', pdf_text, re.IGNORECASE)
        if m_cred:
            t_c = m_cred.group(1).upper()
            val = float(re.sub(r'[^0-9.]', '', t_c.replace(',', '')) or 0)
            cred_m = -val if ('CR' in t_c or '-' in t_c) else val

        # El monto puede aparecer antes o después de la etiqueta "CDT - 2%" según el layout del PDF de Claro
        m_cdt = re.search(r'CDT\s*(?:-|:)\s*2%[^0-9]*([0-9,]+\.[0-9]{2})', pdf_text, re.IGNORECASE)
        if not m_cdt:
            m_cdt = re.search(r'([0-9,]+\.[0-9]{2})\s*CDT\s*(?:-|:)\s*2%', pdf_text, re.IGNORECASE)
        if m_cdt:
            cdt_m = float(m_cdt.group(1).replace(',', ''))

        # ITBIS e ISC impresos literalmente en el PDF (fuente oficial): Claro los calcula por línea
        # con redondeo individual y luego suma, por lo que su total puede diferir en centavos del
        # 18%/10% aplicado de una sola vez sobre el Subtotal consolidado.
        m_itbis = re.search(r'ITBIS\s*(?:-|:)\s*18%[^0-9]*([0-9,]+\.[0-9]{2})', pdf_text, re.IGNORECASE)
        if not m_itbis:
            m_itbis = re.search(r'([0-9,]+\.[0-9]{2})\s*ITBIS\s*(?:-|:)\s*18%', pdf_text, re.IGNORECASE)
        if m_itbis:
            itbis_m = float(m_itbis.group(1).replace(',', ''))

        m_isc = re.search(r'ISC\s*(?:-|:)\s*10%[^0-9]*([0-9,]+\.[0-9]{2})', pdf_text, re.IGNORECASE)
        if not m_isc:
            m_isc = re.search(r'([0-9,]+\.[0-9]{2})\s*ISC\s*(?:-|:)\s*10%', pdf_text, re.IGNORECASE)
        if m_isc:
            isc_m = float(m_isc.group(1).replace(',', ''))

        # Total de Factura Claro impreso literalmente en el PDF (fuente oficial para validar la conciliación)
        total_pdf_m = 0.0
        m_tot_pdf = re.search(r'Total\s+del\s+Mes\s*[\$RD\s]*([0-9,]+\.[0-9]{2})', pdf_text, re.IGNORECASE)
        if not m_tot_pdf:
            m_tot_pdf = re.search(r'(?:Total\s+a\s+Pagar|Total\s+Factura)\s*[\$RD\s]*([0-9,]+\.[0-9]{2})', pdf_text, re.IGNORECASE)
        if m_tot_pdf:
            total_pdf_m = float(m_tot_pdf.group(1).replace(',', ''))

        emp_map = _build_empleado_phone_map(self.env)
        lineas_vals = []
        seen_phones = set()
        # Se registra el body de bienvenida de cada empleado nuevo y se postea en un solo
        # lote al final (ver _message_log_batch más abajo) para no penalizar el tiempo de
        # extracción con un message_post síncrono por cada línea nueva de la factura.
        nuevos_empleados_log = {}

        for line_raw in pdf_text.split('\n'):
            row = extract_claro_phone_row(line_raw)
            if not row:
                continue

            clean_phone = row['phone']
            if clean_phone in ('8092201212', '8092201111'):
                continue
            if clean_phone in seen_phones:
                continue
            seen_phones.add(clean_phone)

            values = row['values']
            if len(values) < 4:
                continue

            ordered_values = values[:7]

            mapped_values = map_claro_columns(ordered_values, source_extractor=pdf_extractor)
            emp = emp_map.get(clean_phone) or emp_map.get(clean_phone[-10:]) or (emp_map.get(clean_phone[-7:]) if len(clean_phone) >= 7 else None)
            if not emp:
                emp_name = f"Empleado Flota {clean_phone}"
                emp = self.env['flota.empleado'].with_context(tracking_disable=True, mail_create_nosubscribe=True).create({
                    'name': emp_name,
                    'numero_flota': clean_phone,
                    'cargo': 'Asignación Automática Claro',
                    'departamento_id': False,
                    'ubicacion_id': False,
                    'estado': 'active',
                    'en_ultima_factura': True,
                    'es_nuevo_auto': True,
                    'notas': f'Nuevo número registrado desde Factura Claro ({self.periodo}). Complete la ficha de empleado.'
                })
                nuevos_empleados_log[emp.id] = _("Empleado registrado automáticamente al aparecer un nuevo número en la factura de Claro (%s): <b>%s</b>.") % (self.periodo, clean_phone)
                emp_map[clean_phone] = emp
                if len(clean_phone) >= 10:
                    emp_map[clean_phone[-10:]] = emp
                if len(clean_phone) >= 7:
                    emp_map[clean_phone[-7:]] = emp

            lineas_vals.append({
                'conciliacion_id': self.id,
                'numero_flota': clean_phone,
                'empleado_id': emp.id,
                'origen_linea': 'pdf',
                'llamadas_roaming_otras_llamadas': mapped_values.get('llamadas_roaming_otras_llamadas', 0.0),
                'otros_servicios_datos': mapped_values.get('otros_servicios_datos', 0.0),
                'uso_local_data_movil': mapped_values.get('uso_local_data_movil', 0.0),
                'monto_roaming': mapped_values.get('monto_roaming', 0.0),
                'financiamiento_equipos': mapped_values.get('financiamiento_equipos', 0.0),
                'otros_cargos_descuentos': mapped_values.get('otros_cargos_descuentos', 0.0),
                'impuestos': mapped_values.get('impuestos', 0.0),
                'total': mapped_values.get('total', 0.0),
            })

        if nuevos_empleados_log:
            self.env['flota.empleado'].browse(list(nuevos_empleados_log.keys()))._message_log_batch(bodies=nuevos_empleados_log)

        self.write({
            'renta_mensual': renta_m,
            'renta_otros_servicios': renta_o,
            'uso_data_movil': data_m,
            'llamadas_roaming': roam_m,
            'otros_cargos_creditos': cred_m,
            'cdt_monto': cdt_m,
            'itbis_pdf_extraido': itbis_m,
            'isc_pdf_extraido': isc_m,
            'total_factura_pdf': total_pdf_m,
        })

        # Solo se eliminan las líneas provenientes de una extracción previa del PDF.
        # Las líneas agregadas manualmente por un usuario (ej. un empleado registrado en
        # otra factura que también debe reflejarse aquí) se conservan intactas.
        self.linea_ids.filtered(lambda l: l.origen_linea == 'pdf').unlink()
        if lineas_vals:
            self.env['flota.factura.linea'].create(lineas_vals)

        self.action_generar_resumen_departamentos()
        # Se mantiene en Borrador para que el usuario valide la información extraída antes de conciliar.
        self.write({'estado': 'draft'})

        self.message_post(body=_("<b>Factura PDF procesada NATIVAMENTE en Odoo:</b><br/>Líneas encontradas: %s | Gran Total: RD$%s<br/>Revise y valide los datos extraídos antes de confirmar la conciliación.") % (len(seen_phones), self.total_mes))

        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': _('Extracción y Conciliación Exitosa'),
                'message': _('Se procesó la factura PDF nativamente. Se extrajeron %s líneas telefónicas.') % len(seen_phones),
                'type': 'success',
                'sticky': False,
                'next': {
                    'type': 'ir.actions.act_window',
                    'res_model': self._name,
                    'res_id': self.id,
                    'views': [(False, 'form')],
                    'target': 'main',
                }
            }
        }

    def action_exportar_excel(self):
        """Exporta la factura y el consolidado departamental a un libro Excel."""
        self.ensure_one()
        # La validación para exportar usa el mismo criterio de conciliación que estado_cuadre:
        # Total Factura Claro (PDF) vs Total del Mes calculado, NO la suma de líneas de empleados.
        diff = self.total_factura_pdf - self.total_mes
        if abs(diff) >= 0.01:
            raise UserError(_(
                "No se puede exportar la factura a Excel porque existe una diferencia de RD$%s entre el Total de la Factura Claro (RD$%s) y el Total del Mes calculado (RD$%s).\n\n"
                "Por favor, revise los valores de Renta, Otros cargos/créditos y CDT capturados de la factura para que ambos montos coincidan antes de exportar el documento."
            ) % (f"{abs(diff):,.2f}", f"{self.total_factura_pdf:,.2f}", f"{self.total_mes:,.2f}"))

        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Conciliación Factura Claro"
        
        header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Calibri", size=14, bold=True, color="1F2937")
        bold_font = Font(name="Calibri", size=11, bold=True)
        summary_title_fill = PatternFill(start_color="374151", end_color="374151", fill_type="solid")
        total_fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")

        ws["A1"] = f"CONCILIACIÓN FINANCIERA FACTURA CLARO — {self.name}"
        ws["A1"].font = title_font
        ws["A2"] = f"Periodo: {self.periodo} | Fecha: {self.fecha_factura} | Estado: {self.estado.upper()}"
        ws["A2"].font = Font(italic=True, color="4B5563")

        # Columnas exactas en el mismo orden que en la vista: EMPLEADO, NÚMERO FLOTA, CARGO, DEPARTAMENTO, ORIGEN, TOTAL LÍNEA
        headers_emp = ["EMPLEADO", "NÚMERO FLOTA", "CARGO", "DEPARTAMENTO", "ORIGEN", "TOTAL LÍNEA (RD$)"]
        for col_num, h in enumerate(headers_emp, 1):
            cell = ws.cell(row=4, column=col_num, value=h)
            cell.fill = header_fill
            cell.font = header_font

        origen_labels = {'pdf': 'Extraído del PDF', 'manual': 'Agregado Manualmente'}
        current_row = 5
        for l in self.linea_ids:
            ws.cell(row=current_row, column=1, value=l.empleado_id.name if l.empleado_id else 'NO REGISTRADO')
            ws.cell(row=current_row, column=2, value=l.numero_flota)
            ws.cell(row=current_row, column=3, value=l.cargo or '')
            ws.cell(row=current_row, column=4, value=l.departamento_id.name if l.departamento_id else 'N/A')
            ws.cell(row=current_row, column=5, value=origen_labels.get(l.origen_linea, l.origen_linea))
            c_tot = ws.cell(row=current_row, column=6, value=l.total_linea)
            c_tot.number_format = '"RD$"#,##0.00'
            current_row += 1

        ws_dept = wb.create_sheet("Consolidado por Departamento")
        dept_headers = [
            "DEPARTAMENTO",
            "CEDI / UBICACIÓN",
            "EMPLEADOS",
            "SUBTOTAL (RD$)",
            "TOTAL DEPTO (RD$)",
            "% DEL TOTAL GENERAL",
        ]
        for col_num, header in enumerate(dept_headers, 1):
            cell = ws_dept.cell(row=1, column=col_num, value=header)
            cell.fill = header_fill
            cell.font = header_font

        for row_num, summary in enumerate(self.resumen_depto_ids, 2):
            ws_dept.cell(row=row_num, column=1, value=summary.departamento_id.name or "Sin departamento")
            ws_dept.cell(row=row_num, column=2, value=summary.ubicacion_id.name or "Sin ubicación")
            ws_dept.cell(row=row_num, column=3, value=summary.cantidad_empleados)
            subtotal_cell = ws_dept.cell(row=row_num, column=4, value=summary.monto_subtotal)
            total_cell = ws_dept.cell(row=row_num, column=5, value=summary.monto_total)
            subtotal_cell.number_format = '"RD$"#,##0.00'
            total_cell.number_format = '"RD$"#,##0.00'
            pct_cell = ws_dept.cell(row=row_num, column=6, value=summary.porcentaje_gasto)
            pct_cell.number_format = '0.00%'

        for column_cells in ws_dept.columns:
            max_len = max(len(str(cell.value or '')) for cell in column_cells)
            column_letter = openpyxl.utils.get_column_letter(column_cells[0].column)
            ws_dept.column_dimensions[column_letter].width = max(max_len + 3, 16)

        # Hoja dedicada de auditoría: solo las líneas agregadas manualmente (empleados de
        # otras facturas u otros ajustes que un usuario incluyó a mano en esta conciliación),
        # separadas de los datos extraídos automáticamente del PDF de Claro.
        lineas_manuales = self.linea_ids.filtered(lambda l: l.origen_linea == 'manual')
        if lineas_manuales:
            ws_manual = wb.create_sheet("Registros Agregados Manualmente")
            manual_headers = ["EMPLEADO", "NÚMERO FLOTA", "CARGO", "DEPARTAMENTO", "TOTAL LÍNEA (RD$)"]
            for col_num, header in enumerate(manual_headers, 1):
                cell = ws_manual.cell(row=1, column=col_num, value=header)
                cell.fill = header_fill
                cell.font = header_font
            for row_num, l in enumerate(lineas_manuales, 2):
                ws_manual.cell(row=row_num, column=1, value=l.empleado_id.name if l.empleado_id else 'NO REGISTRADO')
                ws_manual.cell(row=row_num, column=2, value=l.numero_flota)
                ws_manual.cell(row=row_num, column=3, value=l.cargo or '')
                ws_manual.cell(row=row_num, column=4, value=l.departamento_id.name if l.departamento_id else 'N/A')
                c_tot_manual = ws_manual.cell(row=row_num, column=5, value=l.total_linea)
                c_tot_manual.number_format = '"RD$"#,##0.00'
            for column_cells in ws_manual.columns:
                max_len = max(len(str(cell.value or '')) for cell in column_cells)
                column_letter = openpyxl.utils.get_column_letter(column_cells[0].column)
                ws_manual.column_dimensions[column_letter].width = max(max_len + 3, 16)

        current_row += 2

        ws.cell(row=current_row, column=1, value="RESUMEN GENERAL Y TRIBUTACIÓN FACTURA CLARO").font = bold_font
        ws.cell(row=current_row, column=1).fill = summary_title_fill
        ws.cell(row=current_row, column=1).font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        ws.cell(row=current_row, column=2, value="MONTO (RD$)").fill = summary_title_fill
        ws.cell(row=current_row, column=2).font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        current_row += 1

        rubros = [
            ("Renta Mensual", self.renta_mensual),
            ("Renta Otros Servicios", self.renta_otros_servicios),
            ("Uso Servicios Data Móvil", self.uso_data_movil),
            ("Llamadas Roaming", self.llamadas_roaming),
            ("Otros Cargos, Créditos o Descuentos (CR)", self.otros_cargos_creditos),
            ("SUBTOTAL FACTURA", self.subtotal_factura),
            ("(-) Conceptos Excluidos CDT", self.ajustes_excluidos_cdt),
            ("(=) Base Gravable CDT", self.base_gravable_cdt),
            ("ITBIS - 18%", self.itbis_monto),
            ("CDT - 2%", self.cdt_monto),
            ("ISC - 10%", self.isc_monto),
            ("TOTAL DEL MES", self.total_mes),
        ]

        for name, val in rubros:
            c_name = ws.cell(row=current_row, column=1, value=name)
            c_val = ws.cell(row=current_row, column=2, value=val)
            c_val.number_format = '"RD$"#,##0.00'
            
            if name in ["SUBTOTAL FACTURA", "(=) Base Gravable CDT", "TOTAL DEL MES"]:
                c_name.font = bold_font
                c_val.font = bold_font
            if name == "TOTAL DEL MES":
                c_name.fill = total_fill
                c_val.fill = total_fill

            current_row += 1

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        file_data = base64.b64encode(output.read())
        output.close()

        month_names = (
            "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
            "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre",
        )
        month_name = month_names[self.fecha_factura.month - 1] if self.fecha_factura else (self.periodo or "Sin fecha")
        attachment = self.env['ir.attachment'].create({
            'name': f'Factura Claro {month_name}.xlsx',
            'datas': file_data,
            'res_model': self._name,
            'res_id': self.id,
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        })

        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/{attachment.id}?download=true',
            'target': 'self',
        }

    def action_buscar_lineas_empleado(self):
        """Abre una vista de lista con búsqueda nativa (por nombre de empleado
        o por número de flota) de las líneas de esta factura, sin alterar la
        tabla editable de la pestaña 'Desglose por empleado'."""
        self.ensure_one()
        return {
            'type': 'ir.actions.act_window',
            'name': f'Buscar Empleado / Número - {self.name}',
            'res_model': 'flota.factura.linea',
            'view_mode': 'list,form',
            'views': [
                (self.env.ref('gestion_flota_empleados.view_flota_factura_linea_list').id, 'list'),
                (False, 'form'),
            ],
            'search_view_id': [self.env.ref('gestion_flota_empleados.view_flota_factura_linea_search').id, 'search'],
            'domain': [('conciliacion_id', '=', self.id)],
            'context': {'search_default_conciliacion_id': self.id},
            'target': 'current',
        }


class FlotaFacturaLinea(models.Model):
    _name = 'flota.factura.linea'
    _description = 'Detalle de Factura por Línea / Empleado'
    _order = 'total_linea desc, id asc'

    conciliacion_id = fields.Many2one('flota.factura.conciliacion', string='Factura Conciliación', ondelete='cascade', index=True)
    periodo = fields.Char(string='Periodo', related='conciliacion_id.periodo', store=True, readonly=True)
    fecha_factura = fields.Date(string='Fecha de Factura', related='conciliacion_id.fecha_factura', store=True, readonly=True)
    fecha_subida = fields.Datetime(string='Fecha de Subida', related='conciliacion_id.fecha_subida', store=True, readonly=True)
    numero_flota = fields.Char(string='Número Flota', required=True, index=True)
    origen_linea = fields.Selection([
        ('pdf', 'Extraído del PDF'),
        ('manual', 'Agregado Manualmente'),
    ], string='Origen', default='manual', index=True, help="Indica si la línea proviene de la extracción automática del PDF de Claro o si fue agregada manualmente por un usuario (ej. un empleado de otra factura que también debe reflejarse aquí). Las líneas manuales NO se eliminan al volver a extraer el PDF.")
    
    empleado_id = fields.Many2one('flota.empleado', string='Empleado', ondelete='set null', index=True)
    departamento_id = fields.Many2one('flota.departamento', string='Departamento', related='empleado_id.departamento_id', store=True, readonly=True)
    ubicacion_id = fields.Many2one('flota.ubicacion', string='CEDI / Ubicación', related='empleado_id.ubicacion_id', store=True, readonly=True)
    cargo = fields.Char(string='Cargo', related='empleado_id.cargo', readonly=True)

    llamadas_roaming_otras_llamadas = fields.Monetary(string='Llamadas, roaming y otras llamadas', currency_field='currency_id', default=0.0)
    otros_servicios_datos = fields.Monetary(string='Otros servicios y datos', currency_field='currency_id', default=0.0)
    uso_local_data_movil = fields.Monetary(string='Uso local y data móvil', currency_field='currency_id', default=0.0)
    monto_roaming = fields.Monetary(string='Roaming / LD', currency_field='currency_id', default=0.0)
    financiamiento_equipos = fields.Monetary(string='Financiamiento de equipos', currency_field='currency_id', default=0.0)
    otros_cargos_descuentos = fields.Monetary(string='Otros cargos y descuentos', currency_field='currency_id', default=0.0)
    impuestos = fields.Monetary(string='Impuestos', currency_field='currency_id', default=0.0)
    total = fields.Monetary(
        string='Total',
        currency_field='currency_id',
        default=0.0,
        help='Total de la línea tal como aparece en la factura PDF de Claro.'
    )
    diferencia_pdf = fields.Monetary(
        string='Diferencia PDF',
        compute='_compute_linea_totals',
        store=True,
        currency_field='currency_id',
        help='Diferencia entre el total calculado por Odoo y el total impreso en el PDF.'
    )

    subtotal_linea = fields.Monetary(string='Subtotal Línea', compute='_compute_linea_totals', store=True, currency_field='currency_id')
    itbis_linea = fields.Monetary(string='ITBIS (18%)', compute='_compute_linea_totals', store=True, currency_field='currency_id')
    cdt_linea = fields.Monetary(string='CDT (2%)', compute='_compute_linea_totals', store=True, currency_field='currency_id')
    isc_linea = fields.Monetary(string='ISC (10%)', compute='_compute_linea_totals', store=True, currency_field='currency_id')
    total_linea = fields.Monetary(string='Total Línea (RD$)', compute='_compute_linea_totals', store=True, currency_field='currency_id')

    currency_id = fields.Many2one('res.currency', string='Moneda', related='conciliacion_id.currency_id', store=True)

    estado_linea = fields.Selection([
        ('ok', 'Normal'),
        ('sin_consumo', 'Sin Consumo / RD$0'),
        ('exceso_data_roaming', 'Exceso Data y Roaming'),
        ('exceso_data', 'Exceso Data'),
        ('exceso_roaming', 'Exceso Roaming'),
        ('desconocido', 'Número No Registrado')
    ], string='Estado Línea', compute='_compute_estado_linea', store=True, default='ok')

    @api.model_create_multi
    def create(self, vals_list):
        emp_map = _build_empleado_phone_map(self.env)
        for vals in vals_list:
            if not vals.get('empleado_id') and vals.get('numero_flota'):
                norm = _normalize_phone(vals['numero_flota'])
                emp = emp_map.get(norm) or emp_map.get(norm[-10:])
                if emp:
                    vals['empleado_id'] = emp.id
        return super(FlotaFacturaLinea, self).create(vals_list)

    @api.onchange('numero_flota')
    def _onchange_numero_flota(self):
        """Al capturar/editar el Número Flota manualmente en la línea, busca si ya
        existe un empleado registrado con ese número y autocompleta Empleado (y con
        él, Cargo/Departamento por los campos related) sin esperar a guardar.

        Usa una búsqueda directa acotada por los últimos dígitos (en vez de cargar
        TODOS los empleados a memoria como hace _build_empleado_phone_map, pensado
        para procesar el PDF completo) para responder rápido y no bloquear el
        navegador cuando hay muchos empleados registrados.

        Si no hay coincidencia, no hace nada: el usuario puede escribir el nombre
        en el campo Empleado y, si tampoco existe, crearlo desde ahí (opción
        "Crear y editar..." del propio combo, ya que Cargo y Número Flota son
        obligatorios en flota.empleado).
        """
        if not self.numero_flota:
            return
        norm = _normalize_phone(self.numero_flota)
        if not norm or len(norm) < 7:
            return
        try:
            ultimos = norm[-10:]
            candidatos = self.env['flota.empleado'].with_context(active_test=False).search(
                [('numero_flota', 'ilike', ultimos[-7:])]
            )
            emp = next((c for c in candidatos if _normalize_phone(c.numero_flota) == ultimos), False)
            if emp and self.empleado_id != emp:
                self.empleado_id = emp
        except Exception:
            # Nunca debe bloquear la edición manual de la línea: si la búsqueda
            # falla por cualquier motivo, simplemente se omite el autocompletado.
            _logger.warning("No se pudo autocompletar el empleado para el número %s", self.numero_flota, exc_info=True)

    @api.depends('llamadas_roaming_otras_llamadas', 'otros_servicios_datos', 'uso_local_data_movil', 'monto_roaming', 'financiamiento_equipos', 'otros_cargos_descuentos', 'impuestos', 'total')
    def _compute_linea_totals(self):
        for rec in self:
            sub = rec.llamadas_roaming_otras_llamadas + rec.otros_servicios_datos + rec.uso_local_data_movil + rec.monto_roaming + rec.financiamiento_equipos + rec.otros_cargos_descuentos
            rec.subtotal_linea = sub
            
            base_imponible = (
                rec.llamadas_roaming_otras_llamadas
                + rec.otros_servicios_datos
                + rec.uso_local_data_movil
                + rec.monto_roaming
                + rec.otros_cargos_descuentos
            )
            imp = rec.impuestos if rec.impuestos > 0 else (base_imponible * 0.30)
            rec.itbis_linea = imp * 0.60
            rec.cdt_linea = imp * (1.0 / 15.0)
            rec.isc_linea = imp * (1.0 / 3.0)
            rec.total_linea = sub + imp
            rec.diferencia_pdf = rec.total_linea - rec.total if rec.total else 0.0

    @api.depends('empleado_id', 'total_linea', 'uso_local_data_movil', 'monto_roaming')
    def _compute_estado_linea(self):
        for rec in self:
            if not rec.empleado_id:
                rec.estado_linea = 'desconocido'
            elif all(abs(value) < 0.01 for value in (
                rec.llamadas_roaming_otras_llamadas,
                rec.otros_servicios_datos,
                rec.uso_local_data_movil,
                rec.monto_roaming,
                rec.financiamiento_equipos,
                rec.otros_cargos_descuentos,
                rec.impuestos,
            )):
                rec.estado_linea = 'sin_consumo'
            elif rec.monto_roaming > 0.01 and rec.uso_local_data_movil > 0.01:
                rec.estado_linea = 'exceso_data_roaming'
            elif rec.monto_roaming > 0.01:
                rec.estado_linea = 'exceso_roaming'
            elif rec.uso_local_data_movil > 0.01:
                rec.estado_linea = 'exceso_data'
            else:
                rec.estado_linea = 'ok'


class FlotaFacturaDepartamentoResumen(models.Model):
    _name = 'flota.factura.departamento.resumen'
    _description = 'Resumen Consolidado de Factura por Departamento'
    _order = 'monto_total desc'

    def _auto_init(self):
        self.env.cr.execute("""
            ALTER TABLE flota_factura_departamento_resumen 
            ADD COLUMN IF NOT EXISTS currency_id INTEGER;
        """)
        return super()._auto_init()

    conciliacion_id = fields.Many2one('flota.factura.conciliacion', string='Factura Conciliación', ondelete='cascade', index=True)
    departamento_id = fields.Many2one('flota.departamento', string='Departamento')
    ubicacion_id = fields.Many2one('flota.ubicacion', string='CEDI / Ubicación')
    cantidad_empleados = fields.Integer(string='Empleados')
    monto_subtotal = fields.Monetary(string='Subtotal Depto', currency_field='currency_id')
    monto_total = fields.Monetary(string='Total Depto (RD$)', currency_field='currency_id')
    porcentaje_gasto = fields.Float(string='% del Total General', digits=(5, 2))
    currency_id = fields.Many2one('res.currency', string='Moneda', related='conciliacion_id.currency_id', store=True, readonly=True)


class FlotaFacturaConcepto(models.Model):
    _name = 'flota.factura.concepto'
    _description = 'Concepto o Ajuste de Factura (Grava / Excluido de CDT)'
    _order = 'id asc'

    conciliacion_id = fields.Many2one('flota.factura.conciliacion', string='Factura Conciliación', ondelete='cascade', index=True)
    nombre_concepto = fields.Char(string='Concepto / Detalle', required=True)
    monto = fields.Monetary(string='Monto (RD$)', currency_field='currency_id', default=0.0)
    currency_id = fields.Many2one('res.currency', related='conciliacion_id.currency_id', store=True, readonly=True)
    grava_itbis = fields.Boolean(string='Grava ITBIS (18%)', default=True)
    grava_cdt = fields.Boolean(string='Grava CDT (2%)', default=True)
    grava_isc = fields.Boolean(string='Grava ISC (10%)', default=True)
    notas = fields.Char(string='Observaciones')

    @api.onchange('nombre_concepto')
    def _onchange_nombre_concepto(self):
        if self.nombre_concepto:
            nombre = self.nombre_concepto.lower()
            keywords_excluidos_cdt = ['pago atrasado', 'mora', 'recargo', 'multa', 'interes', 'atraso', 'cargo por pago']
            if any(kw in nombre for kw in keywords_excluidos_cdt):
                self.grava_cdt = False
                self.notas = _("Excluido automáticamente de CDT (Concepto no gravado con 2% CDT)")
